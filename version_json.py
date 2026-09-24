"""Snapshot a JSON directory or URL/ID list and persist semantic changes in SQLite."""
from __future__ import annotations

import argparse
import collections
import csv
import datetime as dt
import difflib
import hashlib
import json
import re
import sqlite3
import sys
from contextlib import ExitStack
from pathlib import Path
from urllib.parse import quote, urlparse

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from storage import StagedRecords, PreviousRecords, put_blob, compact_storage

ROOT = Path(__file__).resolve().parent
MISSING = object()


def config(path: Path) -> dict:
    default = json.loads((ROOT / 'config.json').read_text(encoding='utf-8'))
    if not path.exists():
        path.write_text(json.dumps(default, indent=2) + '\n', encoding='utf-8')
        print(f'Created {path}')
    return {**default, **json.loads(path.read_text(encoding='utf-8'))}


def canonical(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(',', ':'))


def digest(value):
    return hashlib.sha256(canonical(value).encode('utf-8')).hexdigest()


def path_matches(path: str, patterns: list[str]) -> bool:
    normalized = re.sub(r'\[[^]]*\]', '[]', path)
    return path in patterns or normalized in patterns


def normalize(value, cfg, path='$'):
    if isinstance(value, dict):
        return {k: normalize(v, cfg, f'{path}.{k}') for k, v in value.items()
                if k not in cfg['ignore_fields'] and not (path == '$' and k in cfg.get('updated_at_fields', []))
                and not path_matches(f'{path}.{k}', cfg['ignore_paths'])}
    if isinstance(value, list):
        result = [normalize(v, cfg, f'{path}[]') for v in value]
        if path_matches(path, cfg['unordered_arrays']):
            result.sort(key=canonical)
        return result
    return value


def identity(data, filename, cfg):
    if isinstance(data, dict):
        for field in cfg['id_fields']:
            value = data.get(field)
            if isinstance(value, (str, int)) and not isinstance(value, bool) and str(value).strip():
                return f'{field}:{value}'
    return f'file:{filename}'


def reported_update(data, cfg):
    if isinstance(data, dict):
        for field in cfg.get('updated_at_fields', ['updated_at']):
            value = data.get(field)
            if isinstance(value, (str, int, float)) and not isinstance(value, bool) and str(value).strip():
                return str(value).strip()
    return None


def parse_update(value):
    if value is None:
        return None
    try:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
            parsed = dt.datetime.combine(dt.date.fromisoformat(value), dt.time(12), dt.timezone.utc)
        elif re.fullmatch(r'\d{10}(?:\.\d+)?', value):
            parsed = dt.datetime.fromtimestamp(float(value), dt.timezone.utc)
        else:
            parsed = dt.datetime.fromisoformat(value.replace('Z', '+00:00'))
            if parsed.tzinfo is None:
                parsed = parsed.replace(tzinfo=dt.timezone.utc)
        return parsed.astimezone(dt.timezone.utc)
    except (ValueError, OverflowError):
        return None


def change_time(kind, reported, previous_time, snapshot_time):
    if kind == 'unchanged':
        return None, None
    if kind == 'modified' and previous_time:
        candidate = parse_update(reported)
        if candidate and previous_time < candidate <= snapshot_time:
            return candidate.isoformat(timespec='seconds'), 'record timestamp'
    return snapshot_time.isoformat(timespec='seconds'), 'snapshot observation'


def leaves(value, path='$'):
    if isinstance(value, (dict, list)) and not value:
        yield path, value
    elif isinstance(value, dict):
        for key, child in value.items():
            yield from leaves(child, f'{path}.{key}')
    elif isinstance(value, list):
        for i, child in enumerate(value):
            yield from leaves(child, f'{path}[{i}]')
    else:
        yield path, value


def count_fields(value):
    return max(1, sum(1 for _ in leaves(value)))


def word_edits(old, new):
    a, b = re.findall(r"\b[\w'-]+\b", old), re.findall(r"\b[\w'-]+\b", new)
    return sum(max(i2-i1, j2-j1) for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b).get_opcodes() if tag != 'equal')


def array_key_fields(spec):
    """Return a validated list of fields from a string or composite key spec."""
    fields = [spec] if isinstance(spec, str) else spec
    if not isinstance(fields, list) or not fields or not all(isinstance(field, str) and field for field in fields):
        raise ValueError(f'array_keys values must be a field name or non-empty field-name list; got {spec!r}')
    return fields


def array_identity(item, fields):
    return tuple(canonical(item[field]) for field in fields)


def array_path(path, fields, identity):
    # Percent encoding keeps brackets and commas in source values from making
    # the human-readable JSON path ambiguous.
    parts = [f'{field}={quote(value, safe="")}' for field, value in zip(fields, identity)]
    return f'{path}[{",".join(parts)}]'


def changes(old, new, cfg, path='$'):
    if old is not MISSING and new is not MISSING and canonical(old) == canonical(new):
        return
    if isinstance(old, dict) and isinstance(new, dict):
        for key in sorted(old.keys() | new.keys()):
            yield from changes(old.get(key, MISSING), new.get(key, MISSING), cfg, f'{path}.{key}')
    elif isinstance(old, list) and isinstance(new, list):
        key_spec = cfg.get('array_keys', {}).get(path)
        if key_spec:
            keys = array_key_fields(key_spec)
        else:
            keys = []
        if keys and all(isinstance(x, dict) and all(key in x for key in keys) for x in old + new):
            a = {array_identity(x, keys): x for x in old}
            b = {array_identity(x, keys): x for x in new}
            if len(a) == len(old) and len(b) == len(new):
                for name in sorted(a.keys() | b.keys()):
                    yield from changes(a.get(name, MISSING), b.get(name, MISSING), cfg, array_path(path, keys, name))
                return
        for i in range(max(len(old), len(new))):
            yield from changes(old[i] if i < len(old) else MISSING,
                               new[i] if i < len(new) else MISSING, cfg, f'{path}[{i}]')
    else:
        # An object/array addition is expanded so new schema fields can be recognised.
        if old is MISSING and isinstance(new, (dict, list)):
            for leaf, val in leaves(new, path):
                yield leaf, MISSING, val
        elif new is MISSING and isinstance(old, (dict, list)):
            for leaf, val in leaves(old, path):
                yield leaf, val, MISSING
        else:
            yield path, old, new


def database(path):
    db = sqlite3.connect(path)
    db.execute('PRAGMA foreign_keys=ON')
    db.executescript('''CREATE TABLE IF NOT EXISTS snapshots (
        id INTEGER PRIMARY KEY, created_at TEXT NOT NULL, previous_id INTEGER,
        source TEXT NOT NULL, note TEXT, config_json TEXT NOT NULL,
        total INTEGER, added INTEGER, removed INTEGER, modified INTEGER, unchanged INTEGER, schema_only INTEGER);
    CREATE TABLE IF NOT EXISTS records (
        snapshot_id INTEGER, record_id TEXT, filename TEXT, hash TEXT, json_text TEXT,
        PRIMARY KEY(snapshot_id, record_id));
    CREATE TABLE IF NOT EXISTS record_changes (
        snapshot_id INTEGER, record_id TEXT, filename TEXT, previous_filename TEXT,
        change_type TEXT, fields_changed INTEGER, fields_total INTEGER, change_ratio REAL, weighted_ratio REAL,
        PRIMARY KEY(snapshot_id, record_id));
    CREATE TABLE IF NOT EXISTS field_changes (
        snapshot_id INTEGER, record_id TEXT, json_path TEXT, old_value TEXT, new_value TEXT,
        schema_only INTEGER, word_edits INTEGER, PRIMARY KEY(snapshot_id, record_id, json_path));
    CREATE INDEX IF NOT EXISTS ix_changes_record ON record_changes(record_id);
    CREATE INDEX IF NOT EXISTS ix_fields_path ON field_changes(json_path);''')
    db.executescript('''CREATE TABLE IF NOT EXISTS record_blobs (
        hash TEXT PRIMARY KEY, payload BLOB NOT NULL, raw_bytes INTEGER NOT NULL,
        classification_json TEXT NOT NULL);
    CREATE INDEX IF NOT EXISTS ix_fields_record ON field_changes(record_id, snapshot_id);
    CREATE INDEX IF NOT EXISTS ix_records_hash ON records(hash);''')
    columns = {row[1] for row in db.execute('PRAGMA table_info(record_changes)')}
    for name in ('event_at', 'event_basis', 'reported_updated_at'):
        if name not in columns:
            db.execute(f'ALTER TABLE record_changes ADD COLUMN {name} TEXT')
    db.execute('''UPDATE record_changes SET event_at=(SELECT created_at FROM snapshots
        WHERE snapshots.id=record_changes.snapshot_id), event_basis='snapshot observation (legacy)'
        WHERE event_at IS NULL AND change_type IN ('added','removed','modified')''')
    db.commit()
    return db


def scan(source, cfg, result=None):
    result = {} if result is None else result
    errors = []
    files = sorted(source.rglob('*.json'))
    for number, file in enumerate(files, 1):
        relative = file.relative_to(source).as_posix()
        try:
            raw = json.loads(file.read_text(encoding='utf-8-sig'))
            record_id = identity(raw, relative, cfg)
            if record_id in result:
                raise ValueError(f'duplicate record ID {record_id} (also {result[record_id][0]})')
            value = normalize(raw, cfg)
            result[record_id] = (relative, value, digest(value), reported_update(raw, cfg))
        except (ValueError, UnicodeError, OSError) as exc:
            errors.append(f'{relative}: {exc}')
        if number % 1000 == 0:
            print(f'Read {number}/{len(files)} files', flush=True)
    return result, errors


def selected_column(headers, column):
    if column:
        if column not in headers:
            raise ValueError(f'List needs a {column!r} column; found {headers}')
        return column
    for candidate in ('modelID', 'model_id', 'id', 'record_id', 'url', 'URL'):
        if candidate in headers:
            return candidate
    raise ValueError(f'Cannot find a modelID, id, or url column; found {headers}')


def source_entries(source: Path, column: str | None):
    """Read URLs/IDs from text, CSV, Excel, or a JSON array/object list."""
    suffix = source.suffix.lower()
    if suffix == '.txt':
        values = [(str(i), line.strip()) for i, line in enumerate(source.read_text(encoding='utf-8-sig').splitlines(), 1)
                  if line.strip() and not line.lstrip().startswith('#')]
    elif suffix == '.csv':
        with source.open(newline='', encoding='utf-8-sig') as stream:
            reader = csv.DictReader(stream)
            field = selected_column(reader.fieldnames or [], column)
            values = [(str(i), str(row.get(field) or '').strip()) for i, row in enumerate(reader, 2)
                      if str(row.get(field) or '').strip()]
    elif suffix == '.xlsx':
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError('Excel input requires: pip install -r requirements.txt') from exc
        workbook = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(x).strip() if x is not None else '' for x in next(rows, ())]
            index = headers.index(selected_column(headers, column))
            values = []
            for number, row in enumerate(rows, 2):
                raw = row[index] if index < len(row) else None
                if raw is None or str(raw).strip() == '':
                    continue
                if isinstance(raw, float) and raw.is_integer():
                    raw = int(raw)
                values.append((str(number), str(raw).strip()))
        finally:
            workbook.close()
    elif suffix == '.json':
        data = json.loads(source.read_text(encoding='utf-8-sig'))
        if isinstance(data, dict):
            key = column if column in data else next((name for name in ('modelIDs', 'model_ids', 'ids', 'urls', 'records') if name in data), None)
            if key is None or key not in data:
                raise ValueError('JSON list needs modelIDs, ids, urls, or records array (or --column)')
            data = data[key]
        if not isinstance(data, list):
            raise ValueError('JSON list input must contain an array')
        values = []
        for number, item in enumerate(data, 1):
            if isinstance(item, dict):
                field = selected_column(list(item), column)
                item = item[field]
            if isinstance(item, bool) or not isinstance(item, (str, int, float)):
                raise ValueError(f'JSON list item {number} must be a URL or modelID')
            if isinstance(item, float) and item.is_integer():
                item = int(item)
            value = str(item).strip()
            if value:
                values.append((str(number), value))
    else:
        raise ValueError('List input must be .txt, .csv, .xlsx, or .json')
    if not values:
        raise ValueError('Source list has no entries')
    seen = set()
    for line, value in values:
        if value in seen:
            raise ValueError(f'Duplicate list entry {value!r} at row {line}')
        seen.add(value)
    return values


class HTTPSession(requests.Session):
    """Reject plaintext requests, including redirects, before sending them."""

    def send(self, request, **kwargs):
        if urlparse(request.url).scheme != 'https':
            raise ValueError('HTTPS is required, including redirect destinations')
        return super().send(request, **kwargs)


def https_session():
    session = HTTPSession()
    session.headers.update({'Accept': 'application/json', 'User-Agent': 'json-versioner/2.0'})
    retries = Retry(total=3, backoff_factor=.5, allowed_methods={'GET'},
                    status_forcelist=(429, 500, 502, 503, 504))
    session.mount('https://', HTTPAdapter(max_retries=retries))
    return session


def scan_urls(source, cfg, template, column, timeout, prefix=None, result=None, session=None):
    if session is None:
        with https_session() as owned_session:
            return scan_urls(source, cfg, template, column, timeout, prefix, result, owned_session)
    entries = source_entries(source, column)
    result = {} if result is None else result
    errors = []
    if template and '{modelID}' not in template and '{id}' not in template:
        raise ValueError('URL template must contain {modelID} or {id}')
    prefix = cfg.get('url_prefix', '') if prefix is None else prefix
    has_ids = any(urlparse(item).scheme not in ('http', 'https') for _, item in entries)
    if has_ids and not template and not prefix:
        raise ValueError('ID list needs url_prefix in config.json, --url-prefix, or --url-template')
    for number, (line, item) in enumerate(entries, 1):
        supplied_url = urlparse(item).scheme in ('http', 'https')
        if supplied_url:
            url = item
        elif template:
            url = template.replace('{modelID}', quote(item, safe='')).replace('{id}', quote(item, safe=''))
        else:
            url = prefix + quote(item, safe='')
        parsed = urlparse(url)
        if parsed.scheme != 'https' or not parsed.netloc:
            errors.append(f'row {line}: expected an HTTPS URL; configure url_prefix or use --url-template for modelIDs')
            continue
        try:
            with session.get(url, timeout=(min(timeout, 10), timeout), stream=True) as response:
                response.raise_for_status()
                payload = bytearray()
                for chunk in response.iter_content(chunk_size=65536):
                    payload.extend(chunk)
                    if len(payload) > 20_000_000:
                        raise ValueError('response exceeds 20 MB (after decompression)')
            raw = json.loads(payload.decode('utf-8-sig'))
            if not isinstance(raw, dict):
                raise ValueError('expected one JSON object per URL')
            rid = identity(raw, url, cfg)
            if not supplied_url:
                expected = f'{cfg["id_fields"][0]}:{item}'
                if rid.startswith('file:'):
                    rid = expected
                elif rid.split(':', 1)[1] != item:
                    raise ValueError(f'returned ID {rid!r} does not match requested modelID {item!r}')
            if rid in result:
                raise ValueError(f'duplicate record ID {rid} (also {result[rid][0]})')
            value = normalize(raw, cfg)
            result[rid] = (url, value, digest(value), reported_update(raw, cfg))
        except (requests.RequestException, TimeoutError, ValueError, UnicodeError, OSError) as exc:
            errors.append(f'row {line}, {url}: {exc}')
        if number % 100 == 0:
            print(f'Fetched {number}/{len(entries)} URLs', flush=True)
    return result, errors


def schema_paths(pending, comparable_count, cfg):
    # Only additions/removals of null or empty values affecting nearly every
    # comparable record qualify. Data-bearing bulk updates remain data changes.
    occurrences = collections.Counter()
    for diffs in pending:
        for path, old, new in diffs:
            if (old is MISSING and new in (None, '', [], {})) or (new is MISSING and old in (None, '', [], {})):
                occurrences[path] += 1
    minimum = max(2, int(comparable_count * float(cfg['schema_field_fraction']) + 0.999999))
    return {path for path, count in occurrences.items() if count >= minimum}


def weight(path, cfg):
    weights = cfg.get('important_fields', {})
    return float(weights.get(path, weights.get(re.sub(r'\[[^]]*\]', '[]', path), 1)))


def run(args):
    with ExitStack() as resources:
        current = StagedRecords()
        resources.callback(current.close)
        return snapshot(args, current, resources)


def snapshot(args, current, resources):
    source = args.source.resolve()
    if not source.is_dir() and not source.is_file():
        raise ValueError(f'Not a directory or list file: {source}')
    cfg = config(args.config.resolve())
    if cfg['invalid_json'] not in ('abort', 'skip'):
        raise ValueError('invalid_json must be abort or skip')
    history = args.history.resolve()
    if source.is_dir() and (history == source or source in history.parents):
        raise ValueError('History must be outside the source directory')
    if source.is_dir():
        if args.url_template:
            raise ValueError('--url-template requires a .txt, .csv, .xlsx, or .json list')
        current, errors = scan(source, cfg, current)
    else:
        current, errors = scan_urls(source, cfg, args.url_template, args.column, args.timeout,
                                    getattr(args, 'url_prefix', None), current)
    if errors:
        print('\n'.join(errors[:20]), file=sys.stderr)
        if source.is_file() or cfg['invalid_json'] == 'abort':
            raise ValueError(f'{len(errors)} source errors; snapshot aborted')
        print(f'Skipping {len(errors)} invalid files. Skipped files may appear removed.', file=sys.stderr)
    if args.dry_run:
        print(f'Dry run: {len(current)} records; no snapshot written')
        return
    history.mkdir(parents=True, exist_ok=True)
    db = database(history / 'history.sqlite3')
    resources.callback(db.close)
    # Serialize writers before choosing the predecessor, so snapshots cannot fork.
    db.execute('BEGIN IMMEDIATE')
    previous_row = db.execute('SELECT id,created_at,config_json FROM snapshots ORDER BY id DESC LIMIT 1').fetchone()
    previous_id = previous_row[0] if previous_row else None
    previous_time = parse_update(previous_row[1]) if previous_row else None
    if previous_row:
        comparison_keys = ('id_fields', 'ignore_fields', 'ignore_paths', 'updated_at_fields', 'unordered_arrays', 'array_keys')
        old_cfg = json.loads(previous_row[2])
        if any(old_cfg.get(k) != cfg.get(k) for k in comparison_keys):
            raise ValueError('Comparison settings changed; use a new --history directory for a new baseline')
    previous = PreviousRecords(db, previous_id)
    common = current.keys() & previous.keys()
    pending = (changes(previous[rid][1], current[rid][1], cfg) for rid in common
               if current.hashes[rid] != previous.metadata[rid][1])
    schema = schema_paths(pending, len(common), cfg)
    counts = collections.Counter()
    now = dt.datetime.now(dt.timezone.utc)
    timestamp = now.isoformat(timespec='seconds')
    with db:
        cur = db.execute('INSERT INTO snapshots(created_at,previous_id,source,note,config_json) VALUES(?,?,?,?,?)',
                         (timestamp, previous_id, str(source), args.note, canonical(cfg)))
        sid = cur.lastrowid
        for rid, hash_ in current.hashes.items():
            if not db.execute('SELECT 1 FROM record_blobs WHERE hash=?', (hash_,)).fetchone():
                put_blob(db, hash_, canonical(current[rid][1]), cfg)
            db.execute('INSERT INTO records VALUES(?,?,?,?,NULL)', (sid, rid, current.filenames[rid], hash_))
        for rid in sorted(current.keys() | previous.keys()):
            if rid in current and rid in previous.metadata and previous.metadata[rid] == (current.filenames[rid], current.hashes[rid]):
                counts['unchanged'] += 1
                continue
            old, new = previous.get(rid), current.get(rid)
            if old is None:
                kind, diff = 'added', []
            elif new is None:
                kind, diff = 'removed', []
            else:
                diff = list(changes(old[1], new[1], cfg)) if old[2] != new[2] else []
                kind = 'modified' if any(path not in schema for path, _, _ in diff) else 'unchanged'
                if diff and kind == 'unchanged':
                    counts['schema_only'] += 1
            counts[kind] += 1
            if kind == 'unchanged' and not diff and old and new and old[0] == new[0]:
                continue
            meaningful = [(p, a, b) for p, a, b in diff if p not in schema]
            total = max(count_fields(old[1]) if old else 0, count_fields(new[1]) if new else 0, 1)
            weighted_total = max(sum(weight(p, cfg) for p, _ in leaves(old[1])) if old else 0,
                                 sum(weight(p, cfg) for p, _ in leaves(new[1])) if new else 0, 1)
            score = sum(weight(p, cfg) for p, _, _ in meaningful)
            reported = new[3] if new and kind == 'modified' else None
            event_at, event_basis = change_time(kind, reported, previous_time, now)
            db.execute('''INSERT INTO record_changes
                (snapshot_id,record_id,filename,previous_filename,change_type,fields_changed,
                 fields_total,change_ratio,weighted_ratio,event_at,event_basis,reported_updated_at)
                VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''',
                       (sid, rid, new[0] if new else old[0], old[0] if old else None,
                        kind, len(meaningful), total, min(len(meaningful)/total, 1), min(score/weighted_total, 1),
                        event_at, event_basis, reported))
            for path, a, b in diff:
                words = word_edits('' if a is MISSING else a, '' if b is MISSING else b) if (a is MISSING or isinstance(a, str)) and (b is MISSING or isinstance(b, str)) and re.search(r'description', path, re.I) else None
                db.execute('INSERT INTO field_changes VALUES(?,?,?,?,?,?,?)',
                           (sid, rid, path, None if a is MISSING else canonical(a),
                            None if b is MISSING else canonical(b), int(path in schema), words))
        db.execute('UPDATE snapshots SET total=?,added=?,removed=?,modified=?,unchanged=?,schema_only=? WHERE id=?',
                   (len(current), counts['added'], counts['removed'], counts['modified'], counts['unchanged'], counts['schema_only'], sid))
    print(f'Snapshot {sid}: {len(current)} records; {counts["added"]} added, {counts["removed"]} removed, '
          f'{counts["modified"]} modified, {counts["unchanged"]} unchanged, {counts["schema_only"]} schema-only')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', nargs='?', type=Path)
    parser.add_argument('--history', type=Path, default=ROOT / 'json_history')
    parser.add_argument('--config', type=Path, default=ROOT / 'config.json')
    parser.add_argument('--note', default='')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--url-template', help='For ID lists, e.g. https://host/api/models/{modelID}')
    parser.add_argument('--url-prefix', help='Override config url_prefix for modelID entries')
    parser.add_argument('--column', help='CSV/Excel column or JSON object key; auto-detected by default')
    parser.add_argument('--timeout', type=float, default=20, help='Read inactivity timeout in seconds (default: 20); connect timeout capped at 10')
    parser.add_argument('--list', action='store_true', help='Show snapshot history')
    parser.add_argument('--compact-storage', action='store_true', help='Back up, compress/deduplicate legacy records, and reclaim free space')
    args = parser.parse_args()
    try:
        if args.timeout <= 0:
            raise ValueError('--timeout must be positive')
        if args.compact_storage:
            path = args.history / 'history.sqlite3'
            if not path.is_file():
                raise ValueError(f'No database at {path}')
            db = database(path)
            try:
                backup = compact_storage(db, path)
                print(f'Compacted {path}; original backup: {backup}')
            finally:
                db.close()
        elif args.list:
            db = database(args.history / 'history.sqlite3') if args.history.exists() else None
            for row in db.execute('SELECT id,created_at,total,added,removed,modified,schema_only,note FROM snapshots ORDER BY id') if db else []:
                print(*row, sep=' | ')
            if db:
                db.close()
        elif args.source:
            run(args)
        else:
            parser.error('source is required unless --list is used')
    except (ValueError, sqlite3.Error, OSError) as exc:
        parser.exit(1, f'Error: {exc}\n')


if __name__ == '__main__':
    main()
